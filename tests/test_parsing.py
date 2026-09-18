"""文档解析：docx 表格 / PDF 页级提取。

回归背景（docx 表格静默丢失）：`_parse_docx` 原本只用 `doc.paragraphs`，
而它**只产出正文直属的 `<w:p>`**；表格内容是
`<w:tbl> → <w:tr> → <w:tc> → <w:p>` 的嵌套结构，不在其中，于是一整块消失。
这是最坏的失败形态 —— 文档状态 `ready`、`chunk_count > 0`、不报任何错，
只是检索永远查不到表格里的内容。现在按文档顺序遍历段落与表格。

不触网：docx 用 python-docx 现场构造，PDF 用下面的 `_build_pdf` 手搓
（纯标准库，不给测试引入 reportlab 之类的生成器依赖）。
"""
import pytest

from app.config import settings
from app.services.chunking import chunk_sections
from app.services.parsing import parse_file


# ---------------------------------------------------------------- docx 工具


def _docx(tmp_path, name="t.docx"):
    from docx import Document as DocxDocument

    return DocxDocument(), tmp_path / name


def _save(doc, path):
    doc.save(str(path))
    return path


# ---------------------------------------------------------------- PDF 工具


def _wrap_pdf(pages_content: list[str]) -> bytes:
    """把每页的内容流打包成最小合法 PDF（纯标准库，不引第三方生成器依赖）。"""
    objs: list[bytes] = []
    first_page_obj = 3
    n = len(pages_content)
    kids = " ".join(f"{first_page_obj + i * 2} 0 R" for i in range(n))
    font_obj = first_page_obj + n * 2
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    for i, ops in enumerate(pages_content):
        content = ops.encode("latin-1")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> "
            f"/Contents {first_page_obj + i * 2 + 1} 0 R >>".encode()
        )
        objs.append(
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
            + content + b"\nendstream"
        )
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _build_pdf(pages: list[list[str]]) -> bytes:
    """每页给一串文本行；空列表 = 该页无文字（模拟扫描件）。

    字体用内置 Helvetica，编码是 latin-1，**故文本只能是 ASCII** —— 要放中文得
    先嵌入 CJK 字体，那正是本文件的构造器刻意不背的复杂度。
    """
    return _wrap_pdf([
        "BT /F1 12 Tf 72 720 Td 14 TL "
        + " ".join(f"({line}) Tj T*" for line in lines)
        + " ET"
        for lines in pages
    ])


def _build_grid_pdf(rows: list[list[str]], xs=(72, 220, 340), y0=720, dy=20) -> bytes:
    """生成「坐标摆位」的表格 PDF：每格一次 Tm，**同一行的格子共用 y 坐标**。

    这才是真实表格 PDF 的画法（reportlab 等表格库、各种报表系统都这么画）。
    用 `_build_pdf` 造不出表格 —— 那个是每行一次 `T*`，压根没有二维布局。
    """
    ops = "".join(
        f"BT /F1 12 Tf 1 0 0 1 {x} {y0 - r * dy} Tm ({cell}) Tj ET\n"
        for r, row in enumerate(rows)
        for x, cell in zip(xs, row)
    )
    return _wrap_pdf([ops])


# ---------------------------------------------------------------- docx 表格


def test_table_and_paragraphs_keep_document_order(tmp_path):
    """表格夹在段落之间时，必须**在原位置**出现，而不是被挪到末尾或丢掉。"""
    doc, path = _docx(tmp_path)
    doc.add_paragraph("前文")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "C"
    t.cell(1, 1).text = "D"
    doc.add_paragraph("后文")
    _save(doc, path)

    texts = [s.text for s in parse_file(path, "docx")]

    assert texts[0] == "前文"
    assert texts[-1] == "后文"
    assert "A" in texts[1] and "D" in texts[1], f"表格未落在段落之间: {texts}"


def test_table_is_rendered_as_markdown_grid(tmp_path):
    """表格必须保持**二维**：拍平成一维会永久丢失行列对应关系，
    之后就再也答不出"温度那一行的单位是什么"。"""
    doc, path = _docx(tmp_path)
    t = doc.add_table(rows=3, cols=3)
    data = [["参数", "单位", "数值"], ["工作温度", "℃", "25"], ["工作压力", "kPa", "101"]]
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    _save(doc, path)

    sections = parse_file(path, "docx")
    assert len(sections) == 1
    lines = sections[0].text.splitlines()

    assert lines[0] == "| 参数 | 单位 | 数值 |"
    assert set(lines[1].replace(" ", "")) == {"|", "-"}, "缺 Markdown 分隔行"
    assert lines[2] == "| 工作温度 | ℃ | 25 |"
    assert lines[3] == "| 工作压力 | kPa | 101 |"


def test_table_is_marked_in_metadata(tmp_path):
    doc, path = _docx(tmp_path)
    doc.add_paragraph("说明")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "X"
    _save(doc, path)

    by_text = {s.text: s.metadata for s in parse_file(path, "docx")}

    assert by_text["说明"] == {}
    assert by_text["| X |"] == {"has_table": True}


def test_single_row_table_has_no_separator_line(tmp_path):
    """只有一行时插分隔行会造出一个"假表头"，反而误导 —— 故不插。"""
    doc, path = _docx(tmp_path)
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "A"
    _save(doc, path)

    assert parse_file(path, "docx")[0].text == "| A |  |"


def test_table_only_document_is_not_lost(tmp_path):
    """这正是当初的回归现场：整份文档只有一张表 → 旧实现产出 0 个 section
    → 入库时抛「解析后没有可用的文本内容」→ 文档被标 failed。"""
    doc, path = _docx(tmp_path)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "型号"
    t.cell(0, 1).text = "功率"
    t.cell(1, 0).text = "XG-1"
    t.cell(1, 1).text = "15W"
    _save(doc, path)

    sections = parse_file(path, "docx")

    assert len(sections) == 1
    assert "XG-1" in sections[0].text and "15W" in sections[0].text


def test_nested_table_is_flattened_into_its_cell(tmp_path):
    """单元格里的嵌套表格也要取到 —— 否则就是重蹈 `doc.paragraphs` 的覆辙：
    只认一层，更深的结构整块消失。"""
    doc, path = _docx(tmp_path)
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "外层"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "内层甲"
    inner.cell(0, 1).text = "内层乙"
    _save(doc, path)

    text = parse_file(path, "docx")[0].text

    assert "外层" in text
    assert "内层甲" in text and "内层乙" in text


def test_cell_text_containing_pipe_is_escaped(tmp_path):
    """单元格里的 `|` 不转义会把 Markdown 表格切碎，行列随之错乱。"""
    doc, path = _docx(tmp_path)
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "a|b"
    _save(doc, path)

    assert parse_file(path, "docx")[0].text == "| a\\|b |  |"


def test_empty_rows_are_dropped(tmp_path):
    doc, path = _docx(tmp_path)
    t = doc.add_table(rows=3, cols=1)
    t.cell(0, 0).text = "有内容"
    _save(doc, path)

    assert parse_file(path, "docx")[0].text == "| 有内容 |"


def test_empty_document_still_yields_nothing_for_ingestion(tmp_path):
    """空文档必须仍然"解析为空" —— 入库据此抛错并把文档标 failed。
    若这里返回了内容，纯图片文档会静默入库成空 chunk。"""
    doc, path = _docx(tmp_path)
    doc.add_paragraph("   ")
    _save(doc, path)

    sections = parse_file(path, "docx")

    assert sections == [type(sections[0])(text="", metadata={})]
    assert chunk_sections(sections) == [], "空节不该产出 chunk"


def test_table_content_survives_chunking(tmp_path):
    """端到端：解析出来的表格文本要能真正进到 chunk 里（旧实现这里根本没有表）。"""
    doc, path = _docx(tmp_path)
    doc.add_paragraph("设备参数如下。")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "工作温度"
    t.cell(0, 1).text = "25℃"
    t.cell(1, 0).text = "工作压力"
    t.cell(1, 1).text = "101kPa"
    _save(doc, path)

    pieces = chunk_sections(parse_file(path, "docx"), chunk_size=500, overlap=100)
    joined = "\n".join(p.content for p in pieces)

    assert "工作温度" in joined and "25℃" in joined
    assert any(p.metadata.get("has_table") for p in pieces)


def test_large_docx_table_keeps_header_in_every_chunk(tmp_path):
    """端到端：docx 的大表格经「解析 → 分块」后，每一块都要带表头。

    表格行没有句子结束符，原先会被 `_split_at_sentence` 硬切在行中间，且第二块
    起没有表头 —— 只剩一串无列名的值，"XX 型号的额定功率是多少"就只能靠猜列序。
    """
    doc, path = _docx(tmp_path)
    t = doc.add_table(rows=31, cols=3)
    t.cell(0, 0).text = "型号"
    t.cell(0, 1).text = "额定功率"
    t.cell(0, 2).text = "重量"
    for i in range(1, 31):
        t.cell(i, 0).text = f"XG-{1000 + i}"
        t.cell(i, 1).text = f"{i * 15}W"
        t.cell(i, 2).text = f"{i * 0.3:.1f}kg"
    _save(doc, path)

    pieces = chunk_sections(parse_file(path, "docx"), chunk_size=500, overlap=100)

    assert len(pieces) >= 2, "30 行表格应被切开"
    for p in pieces:
        assert p.content.startswith("| 型号 | 额定功率 | 重量 |"), f"缺表头: {p.content[:50]!r}"
        for line in p.content.splitlines():
            assert line.startswith("|") and line.endswith("|"), f"残片: {line!r}"


# ---------------------------------------------------------------- markdown
# 回归背景：`_parse_md` 把标题只写进 `metadata["heading"]`，而该字段全仓库只有写、
# 没有任何读。检索侧（向量 + BM25）只索引 chunk 正文，于是标题文字彻底不可检索 ——
# 实测仓库自己的 README.md：修复前 42 个节的标题文字在 chunk 正文里出现 0 次。
# 这几条用例以前一条都不存在，这正是缺陷能长期全绿的原因。


def test_markdown_heading_enters_section_text(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# 环境变量参考\n\nRERANK_API_KEY 留空则跳过在线重排。\n", encoding="utf-8")

    sections = parse_file(path, "md")

    assert sections[0].text.startswith("# 环境变量参考")
    assert "RERANK_API_KEY" in sections[0].text


def test_every_markdown_heading_reaches_chunk_content(tmp_path):
    """端到端：标题进正文是"能被检索到"的前提，各级标题都要在。"""
    path = tmp_path / "b.md"
    path.write_text(
        "# 部署架构\n后端是 FastAPI。\n\n## 数据库\n用 PostgreSQL 18。\n\n"
        "### 环境变量参考\n至少填 LLM_API_KEY。\n",
        encoding="utf-8",
    )

    blob = "\n".join(
        p.content for p in chunk_sections(parse_file(path, "md"), chunk_size=500, overlap=100)
    )

    for title in ("部署架构", "数据库", "环境变量参考"):
        assert title in blob, f"标题 {title!r} 没能进入 chunk 正文"


def test_markdown_heading_without_body_is_not_dropped(tmp_path):
    """紧邻的两个标题（前一个没有正文）原先被整条丢掉 —— 连 metadata 都没有。"""
    path = tmp_path / "c.md"
    path.write_text("# 第一节\n# 第二节\n正文内容。\n", encoding="utf-8")

    sections = parse_file(path, "md")

    assert [s.metadata["heading"] for s in sections] == ["# 第一节", "# 第二节"]
    assert "第一节" in sections[0].text


def test_markdown_heading_metadata_keeps_original_form(tmp_path):
    """metadata 语义不变（仍是带 `#` 的原始标题行），只是正文里也多了一份。"""
    path = tmp_path / "d.md"
    path.write_text("前言。\n\n## 架构\n内容。\n", encoding="utf-8")

    sections = parse_file(path, "md")

    assert [s.metadata["heading"] for s in sections] == ["", "## 架构"]


# ---------------------------------------------------------------- 代码围栏
# 回归背景：`_parse_md` 逐行 `strip()` 且完全不认围栏（``` / ~~~）。


def test_hash_inside_fenced_block_is_not_a_heading(tmp_path):
    """代码块里的 `# 注释` 不是标题 —— 否则会把节劈成两半，还把 heading 元数据
    污染成一句注释（Python / bash / YAML / nginx 配置里 `#` 开头极常见）。"""
    path = tmp_path / "e.md"
    path.write_text(
        "# 真标题\n\n```python\n# 这是代码里的注释，不是标题\ndef f():\n    return 1\n```\n\n正文。\n",
        encoding="utf-8",
    )

    sections = parse_file(path, "md")

    assert [s.metadata["heading"] for s in sections] == ["# 真标题"]
    assert "# 这是代码里的注释，不是标题" in sections[0].text


def test_fenced_code_keeps_its_indentation(tmp_path):
    """缩进是代码语义的一部分：`    return 1` 压成 `return 1` 就等于改坏了代码。"""
    path = tmp_path / "f.md"
    path.write_text("# 标题\n\n```python\ndef f():\n    return 1\n```\n", encoding="utf-8")

    assert "    return 1" in parse_file(path, "md")[0].text


def test_tilde_fence_is_recognized(tmp_path):
    """`~~~` 也是合法围栏，不能只认反引号。"""
    path = tmp_path / "g.md"
    path.write_text("# 标题\n\n~~~\n# 不是标题\n~~~\n", encoding="utf-8")

    sections = parse_file(path, "md")

    assert [s.metadata["heading"] for s in sections] == ["# 标题"]


def test_longer_closing_fence_is_matched(tmp_path):
    """开栏 3 个反引号、闭栏 4 个也算闭合；内部的 `#` 一行都不该变成标题。"""
    path = tmp_path / "h.md"
    path.write_text("# 标题\n\n```\n# 内部\n````\n\n## 后面的标题\n正文。\n", encoding="utf-8")

    sections = parse_file(path, "md")

    assert [s.metadata["heading"] for s in sections] == ["# 标题", "## 后面的标题"]


# ---------------------------------------------------------------- PDF


def test_pdf_pages_become_sections_with_page_numbers(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(_build_pdf([["page one text"], [], ["page three text"]]))

    sections = parse_file(path, "pdf")

    assert [s.metadata for s in sections] == [{"page": 1}, {"page": 3}]
    assert sections[0].text == "page one text"


def test_pdf_page_without_text_is_skipped_not_emitted_empty(tmp_path):
    """无文字页（扫描件）不能产出空节 —— 空节会在 chunking 里被 continue 掉，
    制造"解析成功但没内容"的假象，掩盖下面那条失败路径。"""
    path = tmp_path / "b.pdf"
    path.write_bytes(_build_pdf([["has text"], []]))

    sections = parse_file(path, "pdf")

    assert len(sections) == 1
    assert all(s.text.strip() for s in sections)


def test_pdf_with_no_text_at_all_yields_nothing(tmp_path):
    """纯扫描件 PDF：解析为空 → 入库抛「解析后没有可用的文本内容」→ 标 failed。
    这是**响亮**的失败，是设计意图，别改成静默入库。"""
    path = tmp_path / "c.pdf"
    path.write_bytes(_build_pdf([[], []]))

    assert parse_file(path, "pdf") == []
    assert chunk_sections(parse_file(path, "pdf")) == []


def test_pdf_table_rows_are_reconstructed(tmp_path):
    """PDF 表格的行结构必须保住 —— 这是 `extraction_mode="layout"` 的收益。

    回归背景：`plain` 模式按文字绘制顺序输出，实测一份真实成绩单 PDF 出来
    **385 行碎片**，`院(系)/部：` 这种标签被拆成 5 行、`20计算机科学与技术1`
    被拆成 3 行。换成 layout 模式后同一文件只有 37 行。本用例用坐标摆位的表格
    PDF 固化这条：同一视觉行的单元格必须落在**同一行**。
    """
    path = tmp_path / "grid.pdf"
    path.write_bytes(_build_grid_pdf([
        ["Parameter", "Unit", "Value"],
        ["Temperature", "C", "25"],
        ["Pressure", "kPa", "101"],
    ]))

    lines = [l for l in parse_file(path, "pdf")[0].text.splitlines() if l.strip()]

    assert len(lines) == 3, f"每行数据应各占一行，实际 {len(lines)} 行: {lines}"
    for cell in ("Parameter", "Unit", "Value"):
        assert cell in lines[0], f"{cell!r} 应与同行的其它单元格在一起: {lines[0]!r}"
    assert "Temperature" in lines[1] and "25" in lines[1]


def test_pdf_table_is_still_not_markdown(tmp_path):
    """**限制的固化用例**：layout 模式还原了「行」，但**列只是空格**。

    两个后果，改动解析时请连同 `CLAUDE.md` / `README` 一起更新：

    1. 列边界没有结构，列间距不齐时两列会粘连 —— 真实成绩单里
       `22公共课/必修课` 就是"总学时"与"类别"两列粘在一起。
    2. 没有 `|`，所以 `chunking._TABLE_ROW` 认不出它：**表格按行切分的逻辑对
       PDF 不生效**，超长 PDF 表格仍会走 `_split_at_sentence` 硬切。

    哪天这里断言到 `|` 了，说明 PDF 表格能转成 Markdown 了 —— 那是重大改进，
    请更新文档与 `test_pdf_table_rows_are_reconstructed` 的措辞，而不是改回来。
    """
    path = tmp_path / "grid2.pdf"
    path.write_bytes(_build_grid_pdf([["Parameter", "Unit", "Value"], ["Temperature", "C", "25"]]))

    text = parse_file(path, "pdf")[0].text

    assert "|" not in text, "出现 | 说明已能还原成 Markdown，是重大改进，请更新文档"
    assert "---" not in text
    # 列分隔只是空格，且至少两个 —— 一个空格会与普通词间空格混同
    assert "Parameter  Unit" in text


def test_squeeze_keeps_two_space_column_separator():
    """layout 模式靠空格表列位：压掉冗余留白，但**保留 2 个空格**当列边界记号。

    压成 1 个空格会让列边界与普通词间空格混同，"A B" 到底是两列还是一个短语
    就分不出来了。
    """
    from app.services.parsing import _squeeze

    assert _squeeze("A" + " " * 40 + "B") == "A  B", "长留白应压成 2 空格"
    assert _squeeze("A  B") == "A  B", "2 空格是列边界记号，必须原样保留"
    assert _squeeze("A B") == "A B", "1 空格是词间空格，不该被扩成 2"
    assert _squeeze("A   B    ") == "A  B", "行尾填充应剥掉"


def test_unsupported_type_raises(tmp_path):
    path = tmp_path / "x.xlsx"
    path.write_text("x", encoding="utf-8")

    with pytest.raises(ValueError, match="不支持的文件类型"):
        parse_file(path, "xlsx")


def test_settings_allow_only_parseable_extensions():
    """白名单与 parse_file 的分支必须一致，否则允许上传却解析不了。"""
    assert settings.allowed_extensions == {"txt", "md", "pdf", "docx"}
